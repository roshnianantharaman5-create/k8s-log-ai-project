from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("Cloud2007-Technical-Walkthrough.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(23, 32, 51)
BLUE = RGBColor(66, 99, 235)
TEAL = RGBColor(24, 169, 153)
ORANGE = RGBColor(240, 140, 70)
PINK = RGBColor(221, 107, 155)
INK = RGBColor(30, 39, 60)
MUTED = RGBColor(101, 112, 138)
PAPER = RGBColor(251, 250, 247)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(235, 239, 255)
PALE_TEAL = RGBColor(231, 250, 245)
PALE_ORANGE = RGBColor(255, 240, 230)
LINE = RGBColor(226, 230, 239)


def shape(slide, kind, x, y, w, h, fill=WHITE, line=LINE, radius=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(1)
    return s


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(.03)
    tf.margin_right = Inches(.03)
    tf.margin_top = Inches(.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_lines(slide, lines, x, y, w, h, size=16, color=MUTED, gap=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
    return box


def base(slide, number):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    text(slide, "CLOUD 2007  ·  KUBERNETES LOG ANALYSIS", .55, .26, 6, .25, 8, BLUE, True)
    text(slide, f"{number:02d}", 12.25, .26, .5, .25, 8, MUTED, True, PP_ALIGN.RIGHT)
    shape(slide, MSO_SHAPE.RECTANGLE, .55, 7.1, 12.2, .012, LINE, LINE)
    text(slide, "kind cluster: cloud2007  ·  project.localhost", .55, 7.17, 6, .18, 7, MUTED)


def title(slide, kicker, heading, subtitle=""):
    text(slide, kicker.upper(), .65, .75, 5.8, .25, 9, BLUE, True)
    text(slide, heading, .65, 1.08, 11.5, .7, 30, NAVY, True)
    if subtitle:
        text(slide, subtitle, .67, 1.88, 11.4, .48, 14, MUTED)


def card(slide, x, y, w, h, heading, body, accent=BLUE, fill=WHITE):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, LINE)
    shape(slide, MSO_SHAPE.OVAL, x + .22, y + .22, .18, .18, accent, accent)
    text(slide, heading, x + .52, y + .15, w - .7, .35, 15, NAVY, True)
    text(slide, body, x + .22, y + .68, w - .44, h - .85, 11, MUTED)


def pill(slide, label, x, y, w, fill=PALE_BLUE, color=BLUE):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, .32, fill, fill)
    text(slide, label, x, y + .055, w, .18, 8, color, True, PP_ALIGN.CENTER)


def flow_box(slide, x, label, detail, color):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.02, 2.08, 1.25, WHITE, LINE)
    shape(slide, MSO_SHAPE.OVAL, x + .77, 2.75, .54, .54, color, color)
    text(slide, label, x + .12, 3.32, 1.84, .27, 14, NAVY, True, PP_ALIGN.CENTER)
    text(slide, detail, x + .15, 3.68, 1.78, .35, 9, MUTED, False, PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2, color=ORANGE):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2.5)
    c.line.end_arrowhead = True


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
base(slide, 1)
shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, NAVY, NAVY)
text(slide, "CLOUD 2007", .72, .7, 3, .3, 11, RGBColor(140, 232, 211), True)
text(slide, "From noisy logs\nto a helpful answer.", .72, 1.45, 7.1, 1.55, 36, WHITE, True)
text(slide, "A technical walkthrough of the Kubernetes observability and AI failure-prediction project.", .75, 3.35, 6.6, .7, 17, RGBColor(205, 214, 235))
pill(slide, "KIND · DOCKER DESKTOP", .75, 4.55, 1.8, RGBColor(47, 63, 94), RGBColor(190, 210, 255))
pill(slide, "LGTM STACK", 2.68, 4.55, 1.35, RGBColor(47, 63, 94), RGBColor(190, 210, 255))
pill(slide, "FASTAPI + STREAMLIT", 4.17, 4.55, 1.75, RGBColor(47, 63, 94), RGBColor(190, 210, 255))
shape(slide, MSO_SHAPE.OVAL, 9.0, 1.45, 2.2, 2.2, BLUE, BLUE)
shape(slide, MSO_SHAPE.OVAL, 10.55, 3.55, 1.25, 1.25, TEAL, TEAL)
shape(slide, MSO_SHAPE.OVAL, 8.25, 4.45, 1.3, 1.3, ORANGE, ORANGE)
text(slide, "APP", 9.0, 2.27, 2.2, .35, 18, WHITE, True, PP_ALIGN.CENTER)
text(slide, "LOGS", 10.55, 4.0, 1.25, .25, 11, WHITE, True, PP_ALIGN.CENTER)
text(slide, "AI", 8.25, 4.92, 1.3, .25, 11, WHITE, True, PP_ALIGN.CENTER)
text(slide, "project.localhost", .75, 6.72, 3, .25, 10, RGBColor(160, 176, 207))

# 2. Why
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 2)
title(slide, "The problem", "A running app leaves clues everywhere.", "The project turns those clues into a small, understandable operational loop.")
card(slide, .7, 2.7, 3.7, 2.7, "The app", "Game 2048 is the demo workload. It gives us a browser-facing service to run, watch, and troubleshoot.", BLUE)
card(slide, 4.82, 2.7, 3.7, 2.7, "The evidence", "Logs, metrics, and traces tell us what happened, how much happened, and where time was spent.", TEAL)
card(slide, 8.94, 2.7, 3.7, 2.7, "The answer", "A FastAPI service reads recent logs and returns a risk estimate with supporting event counts.", ORANGE)
arrow(slide, 4.42, 4.03, 4.78, 4.03); arrow(slide, 8.54, 4.03, 8.9, 4.03)
text(slide, "Goal: reduce the distance between “something feels wrong” and “here is the evidence.”", 1.2, 6.03, 10.9, .45, 18, NAVY, True, PP_ALIGN.CENTER)

# 3. Architecture map
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 3)
title(slide, "Technical component map", "Three namespaces, one connected system.", "Ingress is the front door; services route traffic; workloads do the work.")
for x, name, fill, accent in [(0.8, "game-2048", PALE_BLUE, BLUE), (4.63, "monitoring", PALE_TEAL, TEAL), (8.46, "ai-monitoring", PALE_ORANGE, ORANGE)]:
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.55, 3.35, 3.45, fill, LINE)
    text(slide, name, x + .25, 2.82, 2.8, .3, 16, accent, True)
text(slide, "Ingress", 1.07, 3.45, 1.05, .3, 13, NAVY, True); text(slide, "Service", 2.38, 3.45, 1.05, .3, 13, NAVY, True); text(slide, "Pods", 1.74, 4.25, 1.05, .3, 13, NAVY, True)
arrow(slide, 2.12, 3.76, 2.9, 3.76); arrow(slide, 2.9, 3.94, 2.12, 4.25)
text(slide, "Alloy logs", 4.94, 3.45, 1.25, .3, 13, NAVY, True); text(slide, "Loki", 6.85, 3.45, 1.05, .3, 13, NAVY, True); text(slide, "Grafana", 5.84, 4.35, 1.2, .3, 13, NAVY, True)
arrow(slide, 6.2, 3.76, 6.78, 3.76); arrow(slide, 7.25, 3.96, 6.75, 4.35)
text(slide, "FastAPI", 8.82, 3.45, 1.25, .3, 13, NAVY, True); text(slide, "Streamlit", 10.62, 3.45, 1.3, .3, 13, NAVY, True); text(slide, "ML model", 9.72, 4.35, 1.2, .3, 13, NAVY, True)
arrow(slide, 10.1, 3.76, 10.52, 3.76); arrow(slide, 11.1, 3.96, 10.72, 4.35)
arrow(slide, 3.95, 4.05, 4.55, 4.05, MUTED); arrow(slide, 7.82, 4.05, 8.38, 4.05, MUTED)
text(slide, "Mimir stores metrics  ·  Tempo stores traces  ·  MinIO gives Loki shared object storage", .9, 6.4, 11.5, .3, 12, MUTED, False, PP_ALIGN.CENTER)

# 4. Component responsibilities
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 4)
title(slide, "What each part contributes", "Keep the jobs small and the hand-offs visible.")
rows = [("Game 2048", "demo workload", "A real browser target", BLUE), ("Alloy", "collector", "Reads container logs", TEAL), ("Loki", "log store", "Searches recent events", PINK), ("Grafana", "control room", "Charts and explores signals", ORANGE), ("FastAPI", "prediction service", "Turns log windows into JSON", BLUE), ("Streamlit", "human view", "Explains the latest result", TEAL)]
for i, (name, role, desc, color) in enumerate(rows):
    y = 2.45 + i * .63
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, .85, y, 11.65, .48, WHITE, LINE)
    shape(slide, MSO_SHAPE.OVAL, 1.08, y + .13, .2, .2, color, color)
    text(slide, name, 1.48, y + .1, 2.0, .24, 13, NAVY, True)
    text(slide, role, 3.75, y + .1, 2.1, .24, 11, color, True)
    text(slide, desc, 6.2, y + .1, 5.8, .24, 11, MUTED)

# 5. Observability
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 5)
title(slide, "Observability in plain English", "Three questions, three kinds of signals.")
card(slide, .8, 2.65, 3.7, 2.9, "Logs · what happened?", "A timestamped note: “request timed out”, “service started”, or “health check passed”. Stored in Loki.", BLUE)
card(slide, 4.82, 2.65, 3.7, 2.9, "Metrics · how much?", "Numbers over time: request counts, error rate, CPU, memory, and queue size. Stored in Mimir.", TEAL)
card(slide, 8.84, 2.65, 3.7, 2.9, "Traces · where did time go?", "A request’s journey across services. Tempo helps locate the slow part.", ORANGE)
text(slide, "Grafana is the shared window over all three.", 2, 6.25, 9.3, .35, 18, NAVY, True, PP_ALIGN.CENTER)

# 6. ML pipeline
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 6)
title(slide, "How raw logs become a prediction", "The model is a warning light, not a replacement for an engineer.")
for x, label, detail, color in [(0.65, "Collect", "Loki query_range", BLUE), (3.18, "Normalize", "consistent data frame", TEAL), (5.71, "Vectorize", "TF-IDF + numbers", ORANGE), (8.24, "Classify", "logistic regression", PINK), (10.77, "Serve", "FastAPI JSON", BLUE)]:
    flow_box(slide, x, label, detail, color)
for x in [2.75, 5.28, 7.81, 10.34]: arrow(slide, x, 3.65, x + .35, 3.65)
text(slide, "Training artifacts", 1.1, 5.2, 2.2, .25, 12, NAVY, True)
pill(slide, "artifacts/model.joblib", 3.2, 5.12, 2.0, PALE_BLUE, BLUE)
pill(slide, "metrics.json", 5.45, 5.12, 1.25, PALE_TEAL, TEAL)
pill(slide, "feature_schema.json", 6.9, 5.12, 1.75, PALE_ORANGE, ORANGE)
text(slide, "The live endpoint reports probability, risk level, record count, and top event types.", 1.1, 6.1, 11, .35, 14, MUTED, False, PP_ALIGN.CENTER)

# 7. Samples
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 7)
title(slide, "Samples from the running system", "The same evidence is available to people and tools.")
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, .75, 2.4, 5.85, 3.85, NAVY, NAVY)
text(slide, "LIVE PREDICTION", 1.08, 2.75, 2.5, .25, 10, RGBColor(140, 232, 211), True)
text(slide, '{\n  "predicted_failure": false,\n  "failure_probability": 0.2982,\n  "risk_level": "normal",\n  "records_analyzed": 104\n}', 1.05, 3.25, 4.8, 2.2, 17, RGBColor(220, 231, 255), False, font="Consolas")
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.9, 2.4, 5.7, 3.85, PALE_TEAL, PALE_TEAL)
text(slide, "LOKI / LOGQL", 7.25, 2.75, 2.5, .25, 10, TEAL, True)
text(slide, '{namespace="ai-monitoring"}\n  | json\n  | event_type =~\n    "timeout|http_error"', 7.2, 3.25, 4.8, 1.7, 17, NAVY, False, font="Consolas")
text(slide, "Logs are JSONL; the API exposes Prometheus metrics at /metrics/.", 7.25, 5.45, 4.6, .4, 12, MUTED)

# 8. URLs and test path
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 8)
title(slide, "Try the live cluster", "Everything is routed through the kind ingress controller.")
endpoints = [("project.localhost", "this presentation", BLUE), ("game-2048.localhost", "demo application", TEAL), ("ai-monitoring.localhost/predict/current", "live model answer", ORANGE), ("dashboard.localhost", "Streamlit view", PINK), ("grafana.localhost", "observability control room", BLUE)]
for i, (url, desc, color) in enumerate(endpoints):
    y = 2.35 + i * .7
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.1, y, 11.1, .5, WHITE, LINE)
    shape(slide, MSO_SHAPE.OVAL, 1.35, y + .14, .2, .2, color, color)
    text(slide, url, 1.8, y + .1, 5.4, .25, 13, NAVY, True, font="Consolas")
    text(slide, desc, 7.65, y + .1, 3.9, .25, 12, MUTED)
text(slide, "Quick check", 1.1, 6.05, 1.5, .25, 12, NAVY, True)
text(slide, "kubectl get pods -A  ·  kubectl get ingress -A  ·  Invoke-RestMethod http://ai-monitoring.localhost/predict/current", 2.45, 6.05, 9.7, .35, 11, MUTED, False, font="Consolas")

# 9. Deployment
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 9)
title(slide, "Deployment footprint", "A local, repeatable lab that behaves like a small platform.")
card(slide, .8, 2.55, 3.65, 2.75, "Cluster", "kind cloud2007\n3 nodes\nDocker Desktop\nIngress on localhost", BLUE)
card(slide, 4.85, 2.55, 3.65, 2.75, "Monitoring", "LGTM stack\nAlloy collectors\nLoki → MinIO\nMimir + Tempo", TEAL)
card(slide, 8.9, 2.55, 3.65, 2.75, "Application", "ai-monitoring namespace\nFastAPI + Streamlit\nSynthetic event generator\nModel artifacts", ORANGE)
text(slide, "Source of truth", 1.1, 6.05, 1.7, .25, 12, NAVY, True)
text(slide, "github.com/roshnianantharaman5-create/k8s-log-ai-project", 2.95, 6.05, 8.3, .25, 12, BLUE, True, font="Consolas")

# 10. Close
slide = prs.slides.add_slide(prs.slide_layouts[6]); base(slide, 10)
shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, NAVY, NAVY)
text(slide, "THE TAKEAWAY", .75, .85, 2.3, .25, 10, RGBColor(140, 232, 211), True)
text(slide, "Observe first.\nExplain second.\nAct with evidence.", .75, 1.45, 7.5, 1.45, 34, WHITE, True)
text(slide, "This project connects a demo workload, a telemetry pipeline, and an AI-assisted warning layer in one approachable Kubernetes lab.", .78, 3.55, 6.7, .8, 17, RGBColor(205, 214, 235))
shape(slide, MSO_SHAPE.OVAL, 9.15, 2.0, 2.1, 2.1, BLUE, BLUE)
shape(slide, MSO_SHAPE.OVAL, 10.6, 3.8, 1.1, 1.1, TEAL, TEAL)
shape(slide, MSO_SHAPE.OVAL, 8.3, 4.5, 1.2, 1.2, ORANGE, ORANGE)
text(slide, "APP", 9.15, 2.78, 2.1, .3, 17, WHITE, True, PP_ALIGN.CENTER)
text(slide, "LOGS", 10.6, 4.2, 1.1, .2, 9, WHITE, True, PP_ALIGN.CENTER)
text(slide, "AI", 8.3, 4.96, 1.2, .2, 9, WHITE, True, PP_ALIGN.CENTER)
text(slide, "Questions?", .8, 6.45, 2, .3, 14, RGBColor(160, 176, 207), True)

prs.core_properties.title = "Cloud 2007 - Kubernetes Log Analysis and Observability"
prs.core_properties.subject = "Technical walkthrough of the kind-based LGTM and AI monitoring project"
prs.core_properties.author = "Cloud 2007 project"
prs.save(OUT)
print(OUT)
